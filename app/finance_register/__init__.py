"""Append-only, traceable Excel expense register."""
from __future__ import annotations

import os
import threading
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from app.finance_categorize import CategorizedExpense
from app.finance_extract import ExpenseRecord
from app.shared import audit as audit_mod
from app.shared import kill_switch as kill_switch_mod


HEADERS = (
    "date", "supplier", "amount", "gst", "description", "category",
    "payment_reference", "source_document_id", "filed_location",
)
AuditFn = Callable[..., Any]
_LOCKS: dict[str, threading.RLock] = {}
_LOCKS_GUARD = threading.Lock()


def _lock_for(path: Path) -> threading.RLock:
    key = str(path.resolve()).casefold()
    with _LOCKS_GUARD:
        return _LOCKS.setdefault(key, threading.RLock())


class FinanceRegister:
    def __init__(self, path: Path, *, audit: AuditFn | None = None) -> None:
        self._path = path
        self._audit = audit or audit_mod.log_event

    def append(self, expense: ExpenseRecord | CategorizedExpense, *, filed_location: str | None = None) -> dict[str, Any]:
        kill_switch_mod.assert_allows("finance_register.append")
        record, location = _parts(expense, filed_location)
        if record.status != "extracted":
            raise ValueError("only confirmed extracted expenses may enter the register")
        with _lock_for(self._path):
            workbook = _load_or_create(self._path)
            sheet = workbook.active
            existing = self.find(record.document_id, workbook=workbook)
            if existing is not None:
                self._audit(
                    "finance_register", "append_expense", "document", record.document_id, "success",
                    metadata={"result": "already_registered"},
                )
                return existing
            row = {
                "date": record.date,
                "supplier": record.supplier,
                "amount": record.amount,
                "gst": record.gst,
                "description": record.description,
                "category": record.category,
                "payment_reference": record.payment_reference,
                "source_document_id": record.document_id,
                "filed_location": location,
            }
            sheet.append([row[header] for header in HEADERS])
            _atomic_save(workbook, self._path)
        self._audit(
            "finance_register", "append_expense", "document", record.document_id, "success",
            metadata={"result": "appended"},
        )
        return row

    def find(self, source_document_id: str, *, workbook: Any | None = None) -> dict[str, Any] | None:
        owns = workbook is None
        workbook = workbook or _load_or_create(self._path)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        if "source_document_id" not in headers:
            return None
        ref_index = headers.index("source_document_id")
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if len(values) > ref_index and values[ref_index] == source_document_id:
                return dict(zip(headers, values))
        if owns:
            workbook.close()
        return None


def _parts(expense: ExpenseRecord | CategorizedExpense, location: str | None) -> tuple[ExpenseRecord, str | None]:
    if isinstance(expense, CategorizedExpense):
        return expense.expense, location or expense.filed_location
    return expense, location


def _load_or_create(path: Path) -> Any:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
        if sheet.max_row == 1 and all(cell.value is None for cell in sheet[1]):
            sheet.append(list(HEADERS))
        return workbook
    workbook = Workbook()
    workbook.active.append(list(HEADERS))
    return workbook


def _atomic_save(workbook: Any, path: Path) -> None:
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


__all__ = ["FinanceRegister", "HEADERS"]
